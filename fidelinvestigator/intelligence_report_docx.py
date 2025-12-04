"""
Intelligence Report DOCX Generator
===================================

Generatore di report Intelligence in formato DOCX professionale
seguendo la struttura Dutch OSINT Guy.

Author: FidelinvestigatorAI
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
from typing import Dict, List, Any, Optional
import os

from .intelligence_report_framework import (
    IntelligenceReportGenerator,
    ConfidenceLevel,
    SourceReliability,
    InformationAccuracy,
    KeyJudgment,
    Source,
    EntityOfInterest,
    Hypothesis
)


class IntelligenceReportDOCX:
    """
    Generatore di report Intelligence in formato DOCX
    con formattazione professionale
    """

    # Colori corporate
    COLORS = {
        'primary': RGBColor(0, 51, 102),      # Blu scuro
        'secondary': RGBColor(51, 102, 153),   # Blu medio
        'accent': RGBColor(192, 0, 0),         # Rosso
        'success': RGBColor(0, 128, 0),        # Verde
        'warning': RGBColor(255, 153, 0),      # Arancione
        'light': RGBColor(240, 240, 240),      # Grigio chiaro
        'dark': RGBColor(51, 51, 51),          # Grigio scuro
        'high_confidence': RGBColor(0, 128, 0),
        'moderate_confidence': RGBColor(255, 153, 0),
        'low_confidence': RGBColor(192, 0, 0)
    }

    def __init__(self, report_generator: IntelligenceReportGenerator):
        """
        Inizializza il generatore DOCX

        Args:
            report_generator: IntelligenceReportGenerator configurato
        """
        self.report = report_generator
        self.doc = Document()
        self._setup_styles()
        self._setup_document()

    def _setup_document(self):
        """Configura il documento"""
        # Imposta margini
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)

    def _setup_styles(self):
        """Configura gli stili del documento"""
        styles = self.doc.styles

        # Stile titolo principale
        if 'Intelligence Title' not in [s.name for s in styles]:
            title_style = styles.add_style('Intelligence Title', WD_STYLE_TYPE.PARAGRAPH)
            title_style.font.name = 'Arial'
            title_style.font.size = Pt(24)
            title_style.font.bold = True
            title_style.font.color.rgb = self.COLORS['primary']
            title_style.paragraph_format.space_after = Pt(12)
            title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Stile sezione
        if 'Section Header' not in [s.name for s in styles]:
            section_style = styles.add_style('Section Header', WD_STYLE_TYPE.PARAGRAPH)
            section_style.font.name = 'Arial'
            section_style.font.size = Pt(14)
            section_style.font.bold = True
            section_style.font.color.rgb = self.COLORS['primary']
            section_style.paragraph_format.space_before = Pt(18)
            section_style.paragraph_format.space_after = Pt(6)

        # Stile sottosezione
        if 'Subsection Header' not in [s.name for s in styles]:
            subsection_style = styles.add_style('Subsection Header', WD_STYLE_TYPE.PARAGRAPH)
            subsection_style.font.name = 'Arial'
            subsection_style.font.size = Pt(12)
            subsection_style.font.bold = True
            subsection_style.font.color.rgb = self.COLORS['secondary']
            subsection_style.paragraph_format.space_before = Pt(12)
            subsection_style.paragraph_format.space_after = Pt(6)

        # Stile BLUF box
        if 'BLUF Box' not in [s.name for s in styles]:
            bluf_style = styles.add_style('BLUF Box', WD_STYLE_TYPE.PARAGRAPH)
            bluf_style.font.name = 'Arial'
            bluf_style.font.size = Pt(11)
            bluf_style.paragraph_format.left_indent = Cm(0.5)
            bluf_style.paragraph_format.right_indent = Cm(0.5)

    def _add_shading(self, cell, color: str):
        """Aggiunge sfondo a una cella"""
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), color)
        cell._tc.get_or_add_tcPr().append(shading)

    def _get_confidence_color(self, confidence: ConfidenceLevel) -> RGBColor:
        """Restituisce il colore per il livello di confidenza"""
        colors = {
            ConfidenceLevel.HIGH: self.COLORS['high_confidence'],
            ConfidenceLevel.MODERATE: self.COLORS['moderate_confidence'],
            ConfidenceLevel.LOW: self.COLORS['low_confidence']
        }
        return colors.get(confidence, self.COLORS['dark'])

    def generate_cover_page(self):
        """Genera la pagina di copertina"""
        cover = self.report.generate_cover_page()

        # Classificazione header
        if cover['classification'] != 'UNCLASSIFIED':
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"// {cover['classification']} //")
            run.font.bold = True
            run.font.color.rgb = self.COLORS['accent']
            run.font.size = Pt(14)

        # Spazio
        for _ in range(3):
            self.doc.add_paragraph()

        # Logo placeholder
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("[FIDELINVESTIGATOR AI]")
        run.font.size = Pt(18)
        run.font.color.rgb = self.COLORS['primary']

        # Titolo
        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("INTELLIGENCE REPORT")
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = self.COLORS['primary']

        self.doc.add_paragraph()

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(cover['title'])
        run.font.size = Pt(20)
        run.font.color.rgb = self.COLORS['secondary']

        # Spazio
        for _ in range(4):
            self.doc.add_paragraph()

        # Metadata table
        table = self.doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        metadata = [
            ("Report ID", cover['report_id']),
            ("Date", cover['date']),
            ("Analyst", cover['analyst']),
            ("Organization", cover['organization']),
            ("Distribution", cover['distribution'])
        ]

        for i, (label, value) in enumerate(metadata):
            row = table.rows[i]
            row.cells[0].text = label + ":"
            row.cells[0].paragraphs[0].runs[0].font.bold = True
            row.cells[1].text = value

        # Warning
        if cover['warning']:
            self.doc.add_paragraph()
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(cover['warning'])
            run.font.bold = True
            run.font.color.rgb = self.COLORS['accent']

        # Page break
        self.doc.add_page_break()

    def generate_bluf_section(self):
        """Genera la sezione BLUF"""
        bluf = self.report.bluf_generator

        # Titolo sezione
        p = self.doc.add_paragraph("BOTTOM LINE UP FRONT (BLUF)", style='Section Header')

        # Box BLUF
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        self._add_shading(cell, 'E6F2FF')  # Light blue background

        # Scope e Timeframe
        if bluf.scope:
            p = cell.add_paragraph()
            run = p.add_run("SCOPE: ")
            run.font.bold = True
            p.add_run(bluf.scope)

        if bluf.timeframe:
            p = cell.add_paragraph()
            run = p.add_run("PERIODO: ")
            run.font.bold = True
            p.add_run(bluf.timeframe)

        # Key Judgments
        cell.add_paragraph()
        p = cell.add_paragraph()
        run = p.add_run("KEY JUDGMENTS")
        run.font.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = self.COLORS['primary']

        # Ordina per confidenza
        sorted_judgments = sorted(
            bluf.key_judgments,
            key=lambda x: list(ConfidenceLevel).index(x.confidence)
        )

        for i, kj in enumerate(sorted_judgments, 1):
            p = cell.add_paragraph()

            # Indicatore confidenza
            conf_color = self._get_confidence_color(kj.confidence)
            indicator = "●●●" if kj.confidence == ConfidenceLevel.HIGH else \
                       "●●○" if kj.confidence == ConfidenceLevel.MODERATE else "●○○"

            run = p.add_run(f"{i}. [{kj.confidence.value}] {indicator} ")
            run.font.bold = True
            run.font.color.rgb = conf_color

            p.add_run(kj.statement)

            # Caveats
            if kj.caveats:
                p = cell.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                run = p.add_run("Caveats: ")
                run.font.italic = True
                run.font.size = Pt(9)
                p.add_run("; ".join(kj.caveats)).font.size = Pt(9)

            # Alternative view
            if kj.alternative_view:
                p = cell.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                run = p.add_run("Alternative View: ")
                run.font.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = self.COLORS['warning']
                p.add_run(kj.alternative_view).font.size = Pt(9)

        # Legenda confidenza
        cell.add_paragraph()
        p = cell.add_paragraph()
        run = p.add_run("Confidence Levels: ")
        run.font.bold = True
        run.font.size = Pt(9)

        legend = cell.add_paragraph()
        legend.paragraph_format.space_after = Pt(0)
        run = legend.add_run("HIGH (●●●): 80-95% | ")
        run.font.size = Pt(8)
        run.font.color.rgb = self.COLORS['high_confidence']

        run = legend.add_run("MODERATE (●●○): 55-80% | ")
        run.font.size = Pt(8)
        run.font.color.rgb = self.COLORS['moderate_confidence']

        run = legend.add_run("LOW (●○○): 25-55%")
        run.font.size = Pt(8)
        run.font.color.rgb = self.COLORS['low_confidence']

    def generate_toc(self):
        """Genera Table of Contents"""
        self.doc.add_paragraph("TABLE OF CONTENTS", style='Section Header')

        toc_items = [
            "1. Introduction",
            "2. Methodology",
            "3. Findings",
            "4. Timeline",
            "5. Entities of Interest",
            "6. Analysis",
            "7. Assessment",
            "8. Recommendations",
            "9. Source Evaluation",
            "10. Appendices"
        ]

        for item in toc_items:
            p = self.doc.add_paragraph(item)
            p.paragraph_format.left_indent = Cm(1)

        self.doc.add_page_break()

    def generate_introduction(self):
        """Genera sezione Introduction"""
        self.doc.add_paragraph("1. INTRODUCTION", style='Section Header')

        intro = self.report.introduction
        if isinstance(intro, dict):
            # Purpose
            self.doc.add_paragraph("Purpose", style='Subsection Header')
            self.doc.add_paragraph(intro.get('purpose', 'N/A'))

            # Scope
            self.doc.add_paragraph("Scope", style='Subsection Header')
            self.doc.add_paragraph(intro.get('scope', 'N/A'))

            # Limitations
            if intro.get('limitations'):
                self.doc.add_paragraph("Limitations", style='Subsection Header')
                for lim in intro['limitations']:
                    p = self.doc.add_paragraph(lim, style='List Bullet')

            # Intelligence Requirements
            if intro.get('intelligence_requirements'):
                self.doc.add_paragraph("Intelligence Requirements", style='Subsection Header')
                for req in intro['intelligence_requirements']:
                    p = self.doc.add_paragraph(req, style='List Bullet')

    def generate_methodology(self):
        """Genera sezione Methodology"""
        self.doc.add_paragraph("2. METHODOLOGY", style='Section Header')

        meth = self.report.methodology
        if isinstance(meth, dict):
            # Collection Methods
            self.doc.add_paragraph("Collection Methods", style='Subsection Header')
            for method in meth.get('collection_methods', []):
                self.doc.add_paragraph(method, style='List Bullet')

            # Analysis Techniques
            self.doc.add_paragraph("Analysis Techniques", style='Subsection Header')
            for tech in meth.get('analysis_techniques', []):
                self.doc.add_paragraph(tech, style='List Bullet')

            # Tools Used
            if meth.get('tools_used'):
                self.doc.add_paragraph("Tools Used", style='Subsection Header')
                for tool in meth['tools_used']:
                    self.doc.add_paragraph(tool, style='List Bullet')

            # Intelligence Cycle
            self.doc.add_paragraph("Intelligence Cycle Alignment", style='Subsection Header')
            p = self.doc.add_paragraph(meth.get('intelligence_cycle_phase', ''))
            p.paragraph_format.left_indent = Cm(0.5)

            # Ethical Considerations
            if meth.get('ethical_considerations'):
                self.doc.add_paragraph("Ethical Considerations", style='Subsection Header')
                for eth in meth['ethical_considerations']:
                    self.doc.add_paragraph(eth, style='List Bullet')

    def generate_findings(self):
        """Genera sezione Findings"""
        self.doc.add_paragraph("3. FINDINGS", style='Section Header')

        for finding in self.report.findings:
            # Finding header
            p = self.doc.add_paragraph()
            run = p.add_run(f"[{finding['id']}] {finding['title']}")
            run.font.bold = True
            run.font.size = Pt(11)

            # Confidence badge
            confidence = ConfidenceLevel[finding['confidence']]
            conf_color = self._get_confidence_color(confidence)

            p = self.doc.add_paragraph()
            run = p.add_run(f"Confidence: {finding['confidence']}")
            run.font.color.rgb = conf_color
            run.font.bold = True
            run.font.size = Pt(10)

            # Category
            if finding.get('category'):
                p.add_run(f" | Category: {finding['category']}")

            # Description
            self.doc.add_paragraph(finding['description'])

            # Evidence
            if finding.get('evidence'):
                p = self.doc.add_paragraph()
                run = p.add_run("Evidence:")
                run.font.bold = True
                for ev in finding['evidence']:
                    self.doc.add_paragraph(ev, style='List Bullet')

            # Sources
            if finding.get('sources'):
                p = self.doc.add_paragraph()
                run = p.add_run("Sources: ")
                run.font.italic = True
                run.font.size = Pt(9)
                sources_text = ", ".join([s['name'] + f" ({s['rating']})" for s in finding['sources']])
                p.add_run(sources_text).font.size = Pt(9)

            self.doc.add_paragraph()  # Spacing

    def generate_timeline(self):
        """Genera sezione Timeline"""
        self.doc.add_paragraph("4. TIMELINE", style='Section Header')

        if not self.report.timeline:
            self.doc.add_paragraph("No timeline events recorded.")
            return

        # Crea tabella timeline
        table = self.doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'

        # Header
        header_cells = table.rows[0].cells
        headers = ['Date', 'Event', 'Significance']
        for i, header in enumerate(headers):
            header_cells[i].text = header
            header_cells[i].paragraphs[0].runs[0].font.bold = True
            self._add_shading(header_cells[i], 'D9E2F3')

        # Eventi
        for event in self.report.timeline:
            row = table.add_row()
            row.cells[0].text = event['date'][:10]  # Solo data
            row.cells[1].text = event['event']
            row.cells[2].text = event['significance'].upper()

            # Colore per significance
            sig = event['significance'].lower()
            if sig == 'high':
                self._add_shading(row.cells[2], 'FFCCCC')
            elif sig == 'medium':
                self._add_shading(row.cells[2], 'FFEECC')

    def generate_entities(self):
        """Genera sezione Entities of Interest"""
        self.doc.add_paragraph("5. ENTITIES OF INTEREST", style='Section Header')

        if not self.report.entities:
            self.doc.add_paragraph("No entities of interest identified.")
            return

        for entity in self.report.entities:
            # Entity header
            p = self.doc.add_paragraph()
            run = p.add_run(f"{entity.name}")
            run.font.bold = True
            run.font.size = Pt(11)
            p.add_run(f" ({entity.entity_type})")

            # Risk level
            p = self.doc.add_paragraph()
            run = p.add_run(f"Risk Level: {entity.risk_level.upper()}")
            risk_colors = {
                'high': self.COLORS['accent'],
                'medium': self.COLORS['warning'],
                'low': self.COLORS['success'],
                'unknown': self.COLORS['dark']
            }
            run.font.color.rgb = risk_colors.get(entity.risk_level.lower(), self.COLORS['dark'])
            run.font.bold = True

            # Identifiers
            if entity.identifiers:
                p = self.doc.add_paragraph()
                run = p.add_run("Identifiers:")
                run.font.bold = True
                for key, value in entity.identifiers.items():
                    self.doc.add_paragraph(f"{key}: {value}", style='List Bullet')

            # Relationships
            if entity.relationships:
                p = self.doc.add_paragraph()
                run = p.add_run("Relationships:")
                run.font.bold = True
                for rel in entity.relationships:
                    self.doc.add_paragraph(str(rel), style='List Bullet')

            self.doc.add_paragraph()  # Spacing

    def generate_analysis(self):
        """Genera sezione Analysis con ACH"""
        self.doc.add_paragraph("6. ANALYSIS", style='Section Header')

        # Narrative analysis
        if self.report.analysis:
            self.doc.add_paragraph(self.report.analysis)

        # ACH Analysis
        ach = self.report.ach_analyzer.generate_report()

        if ach['hypotheses']:
            self.doc.add_paragraph("Analysis of Competing Hypotheses (ACH)", style='Subsection Header')

            # Descrizione ACH
            p = self.doc.add_paragraph()
            p.add_run(
                "L'ACH è una tecnica analitica sviluppata dalla CIA per valutare "
                "ipotesi alternative minimizzando i bias cognitivi."
            ).font.italic = True

            # Tabella ipotesi
            table = self.doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'

            headers = ['Hypothesis', 'Description', 'Probability', 'Inconsistency']
            header_cells = table.rows[0].cells
            for i, header in enumerate(headers):
                header_cells[i].text = header
                header_cells[i].paragraphs[0].runs[0].font.bold = True
                self._add_shading(header_cells[i], 'D9E2F3')

            for h in ach['hypotheses']:
                row = table.add_row()
                row.cells[0].text = h['id']
                row.cells[1].text = h['description']
                row.cells[2].text = f"{h['probability']*100:.1f}%"
                row.cells[3].text = f"{h['inconsistency_score']*100:.1f}%"

            # Most likely hypothesis
            if ach['most_likely']:
                self.doc.add_paragraph()
                p = self.doc.add_paragraph()
                run = p.add_run("Most Likely Hypothesis: ")
                run.font.bold = True
                p.add_run(f"{ach['most_likely']['id']} - {ach['most_likely']['description']}")

            # Diagnostic evidence
            if ach['diagnostic_evidence']:
                self.doc.add_paragraph("Diagnostic Evidence", style='Subsection Header')
                for diag in ach['diagnostic_evidence'][:5]:
                    self.doc.add_paragraph(
                        f"• {diag['evidence']} (Diagnostic Value: {diag['diagnostic_value']*100:.0f}%)"
                    )

    def generate_assessment(self):
        """Genera sezione Assessment"""
        self.doc.add_paragraph("7. ASSESSMENT", style='Section Header')

        assessment = self.report.generate_assessment()

        # Summary
        self.doc.add_paragraph("Summary", style='Subsection Header')
        self.doc.add_paragraph(assessment['summary'])

        # Overall confidence
        self.doc.add_paragraph("Overall Confidence", style='Subsection Header')
        conf = assessment['confidence_assessment']
        p = self.doc.add_paragraph()
        run = p.add_run(f"Level: {conf['level']}")
        run.font.bold = True
        if conf['level'] == 'HIGH':
            run.font.color.rgb = self.COLORS['high_confidence']
        elif conf['level'] == 'MODERATE':
            run.font.color.rgb = self.COLORS['moderate_confidence']
        else:
            run.font.color.rgb = self.COLORS['low_confidence']

        self.doc.add_paragraph(conf.get('explanation', ''))

        # Key Judgments (riassunto)
        self.doc.add_paragraph("Key Judgments Summary", style='Subsection Header')
        for kj in assessment['key_judgments']:
            p = self.doc.add_paragraph()
            run = p.add_run(f"[{kj['confidence']}] ")
            run.font.bold = True
            p.add_run(kj['statement'])

        # Gaps and Uncertainties
        if assessment['gaps_and_uncertainties']:
            self.doc.add_paragraph("Gaps and Uncertainties", style='Subsection Header')
            for gap in assessment['gaps_and_uncertainties']:
                self.doc.add_paragraph(gap, style='List Bullet')

        # Alternative Hypotheses
        if assessment['alternative_hypotheses']:
            self.doc.add_paragraph("Alternative Hypotheses", style='Subsection Header')
            for alt in assessment['alternative_hypotheses']:
                p = self.doc.add_paragraph()
                run = p.add_run(f"{alt['id']}: ")
                run.font.bold = True
                p.add_run(alt['description'])

    def generate_recommendations(self):
        """Genera sezione Recommendations"""
        self.doc.add_paragraph("8. RECOMMENDATIONS", style='Section Header')

        if not self.report.recommendations:
            self.doc.add_paragraph("No recommendations at this time.")
            return

        # Ordina per priorità
        priority_order = {'high': 0, 'medium': 1, 'low': 2}
        sorted_recs = sorted(
            self.report.recommendations,
            key=lambda x: priority_order.get(x['priority'].lower(), 3)
        )

        # Tabella raccomandazioni
        table = self.doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'

        headers = ['Priority', 'Recommendation']
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = header
            header_cells[i].paragraphs[0].runs[0].font.bold = True
            self._add_shading(header_cells[i], 'D9E2F3')

        for rec in sorted_recs:
            row = table.add_row()
            row.cells[0].text = rec['priority'].upper()
            row.cells[1].text = rec['text']

            # Colore priorità
            priority = rec['priority'].lower()
            if priority == 'high':
                self._add_shading(row.cells[0], 'FFCCCC')
            elif priority == 'medium':
                self._add_shading(row.cells[0], 'FFEECC')
            else:
                self._add_shading(row.cells[0], 'CCFFCC')

    def generate_source_evaluation(self):
        """Genera sezione Source Evaluation"""
        self.doc.add_paragraph("9. SOURCE EVALUATION", style='Section Header')

        src_report = self.report.source_evaluator.generate_source_report()

        # Summary stats
        self.doc.add_paragraph("Source Statistics", style='Subsection Header')

        p = self.doc.add_paragraph()
        run = p.add_run(f"Total Sources: {src_report['total_sources']}")
        run.font.bold = True

        # Reliability distribution
        self.doc.add_paragraph("Reliability Distribution", style='Subsection Header')

        if src_report['reliability_distribution']:
            table = self.doc.add_table(rows=1, cols=len(src_report['reliability_distribution']) + 1)
            table.style = 'Table Grid'

            header_cells = table.rows[0].cells
            header_cells[0].text = "Rating"
            self._add_shading(header_cells[0], 'D9E2F3')

            for i, (rating, count) in enumerate(src_report['reliability_distribution'].items(), 1):
                if i < len(header_cells):
                    header_cells[i].text = f"{rating}: {count}"

        # Source list
        if src_report['sources']:
            self.doc.add_paragraph("Source Details", style='Subsection Header')

            table = self.doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'

            headers = ['Source', 'Type', 'Rating', 'Archived']
            header_cells = table.rows[0].cells
            for i, header in enumerate(headers):
                header_cells[i].text = header
                header_cells[i].paragraphs[0].runs[0].font.bold = True
                self._add_shading(header_cells[i], 'D9E2F3')

            for src in src_report['sources']:
                row = table.add_row()
                row.cells[0].text = src['name']
                row.cells[1].text = src['type']
                row.cells[2].text = src['rating']
                row.cells[3].text = "Yes" if src['archived'] else "No"

        # Recommendations
        if src_report.get('recommendations'):
            self.doc.add_paragraph("Source Recommendations", style='Subsection Header')
            for rec in src_report['recommendations']:
                self.doc.add_paragraph(rec, style='List Bullet')

    def generate_appendices(self):
        """Genera sezione Appendices"""
        self.doc.add_paragraph("10. APPENDICES", style='Section Header')

        if not self.report.appendices:
            self.doc.add_paragraph("No appendices.")
            return

        for app in self.report.appendices:
            self.doc.add_paragraph(f"{app['id']}: {app['title']}", style='Subsection Header')

            if app['type'] == 'text':
                self.doc.add_paragraph(str(app['content']))
            elif app['type'] == 'data':
                # Format as code block
                p = self.doc.add_paragraph()
                run = p.add_run(str(app['content']))
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
            elif app['type'] == 'table' and isinstance(app['content'], list):
                if app['content'] and isinstance(app['content'][0], dict):
                    headers = list(app['content'][0].keys())
                    table = self.doc.add_table(rows=1, cols=len(headers))
                    table.style = 'Table Grid'

                    header_cells = table.rows[0].cells
                    for i, header in enumerate(headers):
                        header_cells[i].text = header
                        header_cells[i].paragraphs[0].runs[0].font.bold = True

                    for item in app['content']:
                        row = table.add_row()
                        for i, key in enumerate(headers):
                            row.cells[i].text = str(item.get(key, ''))

    def generate_footer(self):
        """Genera footer del report"""
        self.doc.add_paragraph()
        self.doc.add_paragraph("=" * 50)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"END OF REPORT - {self.report.report_id}")
        run.font.bold = True

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Framework: Dutch OSINT Guy Methodology")
        run.font.italic = True

        # Classification footer
        if self.report.classification != 'UNCLASSIFIED':
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"// {self.report.classification} //")
            run.font.bold = True
            run.font.color.rgb = self.COLORS['accent']

    def generate_full_report(self, output_path: str):
        """
        Genera il report DOCX completo

        Args:
            output_path: Percorso del file di output
        """
        # Genera tutte le sezioni
        self.generate_cover_page()
        self.generate_bluf_section()
        self.doc.add_page_break()
        self.generate_toc()
        self.generate_introduction()
        self.generate_methodology()
        self.doc.add_page_break()
        self.generate_findings()
        self.doc.add_page_break()
        self.generate_timeline()
        self.generate_entities()
        self.doc.add_page_break()
        self.generate_analysis()
        self.doc.add_page_break()
        self.generate_assessment()
        self.generate_recommendations()
        self.doc.add_page_break()
        self.generate_source_evaluation()
        self.generate_appendices()
        self.generate_footer()

        # Salva documento
        self.doc.save(output_path)
        return output_path


def generate_intelligence_report_docx(
    report_generator: IntelligenceReportGenerator,
    output_path: str
) -> str:
    """
    Funzione utility per generare report DOCX

    Args:
        report_generator: IntelligenceReportGenerator configurato
        output_path: Percorso del file di output

    Returns:
        Percorso del file generato
    """
    docx_generator = IntelligenceReportDOCX(report_generator)
    return docx_generator.generate_full_report(output_path)
